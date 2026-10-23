"""Seleção das linhas da planilha e marcação de Documentação = SIM."""

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from teladoc_bot.azure_link import WorkItemRef, parse_work_item_link
from teladoc_bot.config import ColumnMap, ConfigError
from teladoc_bot.textutil import normalize

DOCUMENTED_VALUE = "SIM"
_EMPTY = ""
_NAO = "nao"
_CRUD = "crud"
_HARD_CRUD = "hard crud"
_NAME = 1
_US_ID = 4
_CRUD_KIND = 6
_DOCUMENTATION = 7
_AZURE_LINK = 8
_NOTES = 9


@dataclass(frozen=True)
class SheetRow:
    excel_row: int
    screen_name: str
    documentation: str
    status: str
    azure_link: str
    notes: str
    work_item: WorkItemRef | None
    skip_reason: str | None


def load_rows(path: Path, sheet_name: str | None, header_row: int, columns: ColumnMap) -> list[SheetRow]:
    del columns
    if not path.is_file():
        raise ConfigError(f"Planilha não encontrada: {path}")
    workbook = load_workbook(path, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        worksheet = _resolve_sheet(workbook, sheet_name)
        rows: list[SheetRow] = []
        for excel_row in range(header_row + 1, worksheet.max_row + 1):
            item = _read_row(worksheet, excel_row)
            if item is None:
                continue
            rows.append(item)
        return rows
    finally:
        workbook.close()


def mark_documented(
    path: Path,
    sheet_name: str | None,
    header_row: int,
    columns: ColumnMap,
    excel_row: int,
    expected: WorkItemRef,
) -> None:
    del columns, header_row
    workbook = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        worksheet = _resolve_sheet(workbook, sheet_name)
        sheet_id = _sheet_id(worksheet.cell(excel_row, _US_ID).value)
        if sheet_id != expected.id:
            raise ConfigError(
                "A linha da planilha mudou antes da marcação. "
                "A US pode já ter sido atualizada, e Documentação permanece NÃO."
            )
        worksheet.cell(excel_row, _DOCUMENTATION).value = DOCUMENTED_VALUE
        temporary = path.with_suffix(path.suffix + ".tmp")
        workbook.save(temporary)
    finally:
        workbook.close()
    temporary.replace(path)


def _resolve_sheet(workbook, sheet_name: str | None) -> Worksheet:
    if not sheet_name:
        return workbook.active
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    wanted = normalize(sheet_name)
    matches = [name for name in workbook.sheetnames if normalize(name) == wanted]
    if len(matches) == 1:
        return workbook[matches[0]]
    available = ", ".join(workbook.sheetnames)
    raise ConfigError(f"Aba '{sheet_name}' não encontrada. Abas disponíveis: {available}")


def _read_row(worksheet: Worksheet, excel_row: int) -> SheetRow | None:
    screen_name = _text(worksheet.cell(excel_row, _NAME).value)
    sheet_id = _sheet_id(worksheet.cell(excel_row, _US_ID).value)
    status = _text(worksheet.cell(excel_row, _CRUD_KIND).value)
    documentation = _text(worksheet.cell(excel_row, _DOCUMENTATION).value)
    azure_link = _text(worksheet.cell(excel_row, _AZURE_LINK).value)
    notes = _text(worksheet.cell(excel_row, _NOTES).value)
    if not any((screen_name, sheet_id, status, documentation, azure_link, notes)):
        return None
    if not azure_link and sheet_id is not None:
        azure_link = _link_from_sheet_id(sheet_id)
    work_item = parse_work_item_link(azure_link)
    return SheetRow(
        excel_row=excel_row,
        screen_name=screen_name,
        documentation=documentation,
        status=status,
        azure_link=azure_link,
        notes=notes,
        work_item=work_item,
        skip_reason=_skip_reason(screen_name, documentation, status, azure_link, work_item, sheet_id),
    )


def _skip_reason(
    screen_name: str,
    documentation: str,
    status: str,
    azure_link: str,
    work_item: WorkItemRef | None,
    sheet_id: int | None,
) -> str | None:
    normalized_status = normalize(status)
    if normalized_status == _HARD_CRUD:
        return "status HARD CRUD não é processado nesta versão"
    if normalized_status != _CRUD:
        return "status diferente de CRUD"
    if normalize(documentation) != _NAO:
        return 'documentação diferente de "NÃO"'
    if not azure_link.strip():
        return "link calculado ausente"
    if work_item is None:
        return "link inválido de work item do Azure DevOps"
    if sheet_id is None or work_item.id != sheet_id:
        return "divergência entre o ID da US e o link Azure"
    if not screen_name.strip():
        return "nome da tela vazio"
    return None


def _link_from_sheet_id(sheet_id: int) -> str:
    organization = os.environ.get("AZURE_ORGANIZATION", "").strip().rstrip("/")
    project = os.environ.get("AZURE_PROJECT", "").strip().strip("/")
    if not organization or not project:
        return _EMPTY
    return f"{organization}/{project}/_workitems/edit/{sheet_id}"


def _sheet_id(value: object) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        if value.is_integer() and value > 0:
            return int(value)
        return None
    text = str(value).replace("\xa0", " ").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if text.isdigit():
        number = int(text)
        return number if number > 0 else None
    return None


def _text(value: object) -> str:
    if value is None:
        return _EMPTY
    return str(value).replace("\xa0", " ").strip()
