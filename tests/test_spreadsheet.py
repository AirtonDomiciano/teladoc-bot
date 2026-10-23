import unittest
from pathlib import Path

import support  # noqa: F401
from openpyxl import Workbook, load_workbook

from teladoc_bot.azure_link import parse_work_item_link
from teladoc_bot.config import ColumnMap, ConfigError
from support import write_workbook
from teladoc_bot.spreadsheet import load_rows, mark_documented

COLUMNS = ColumnMap(
    screen_name="Nome da tela",
    documentation="documentacao",
    status="Status",
    azure_link="Link Azure",
    notes="Observações",
)


class SpreadsheetTests(unittest.TestCase):
    def test_selection_rules(self) -> None:
        path = Path(self.id().replace(".", "_") + ".xlsx")
        try:
            write_workbook(
                path,
                [
                    ["Clientes", "Não", "CRUD", "https://dev.azure.com/org/proj/_workitems/edit/10", "manter filtro"],
                    ["", "", "", "", ""],
                    ["Árduo", "nao", "HÁRD CRUD", "https://dev.azure.com/org/proj/_workitems/edit/11", ""],
                    ["Sem link", " Não ", " Crud ", "", ""],
                    ["Link ruim", "NÃO", "CRÚD", "https://example.com/item", ""],
                    ["Pronta", "Sim", "CRUD", "https://dev.azure.com/org/proj/_workitems/edit/12", ""],
                    ["Outro status", "Não", "leitura", "https://dev.azure.com/org/proj/_workitems/edit/13", ""],
                ],
                headers=["Nome da tela", "Documentação", "Status", "Link Azure", "Observações"],
            )
            rows = load_rows(path, "telas", 1, COLUMNS)
        finally:
            path.unlink(missing_ok=True)
        reasons = {row.screen_name: row.skip_reason for row in rows}
        self.assertEqual(reasons["Clientes"], None)
        self.assertEqual(rows[0].work_item.id if rows[0].work_item else None, 10)
        self.assertEqual(reasons["Árduo"], "status HARD CRUD não é processado nesta versão")
        self.assertEqual(reasons["Sem link"], "sem link")
        self.assertEqual(reasons["Link ruim"], "link inválido de work item do Azure DevOps")
        self.assertEqual(reasons["Pronta"], 'documentação diferente de "Não"')
        self.assertEqual(reasons["Outro status"], "status diferente de CRUD")
        self.assertNotIn("", reasons)

    def test_hyperlink_and_formula(self) -> None:
        path = Path(self.id().replace(".", "_") + ".xlsx")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Telas"
        worksheet.append(["Nome da tela", "Documentação", "Status", "Link Azure", "Observações"])
        worksheet.append(["Por hiperlink", "Não", "CRUD", "abrir", ""])
        worksheet.cell(2, 4).hyperlink = "https://dev.azure.com/org/proj/_workitems/edit/21"
        worksheet.cell(2, 4).value = "abrir"
        worksheet.append(
            [
                "Por fórmula",
                "Não",
                "CRUD",
                '=HYPERLINK("https://dev.azure.com/org/proj/_workitems/edit/22","US")',
                "",
            ]
        )
        workbook.save(path)
        workbook.close()
        try:
            rows = load_rows(path, None, 1, COLUMNS)
        finally:
            path.unlink(missing_ok=True)
        self.assertEqual([row.work_item.id if row.work_item else None for row in rows], [21, 22])

    def test_mark_sim_only_when_the_link_still_matches(self) -> None:
        path = Path(self.id().replace(".", "_") + ".xlsx")
        write_workbook(
            path,
            [["Clientes", "Não", "CRUD", "https://dev.azure.com/org/proj/_workitems/edit/10", "nota"]],
        )
        ref = parse_work_item_link("https://dev.azure.com/org/proj/_workitems/edit/10")
        other = parse_work_item_link("https://dev.azure.com/org/proj/_workitems/edit/99")
        assert ref is not None and other is not None
        try:
            with self.assertRaises(ConfigError):
                mark_documented(path, "Telas", 1, COLUMNS, 2, other)
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["B2"].value, "Não")
            workbook.close()
            mark_documented(path, "Telas", 1, COLUMNS, 2, ref)
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["B2"].value, "Sim")
            self.assertEqual(workbook.active["E2"].value, "nota")
            self.assertEqual(workbook.active["A2"].value, "Clientes")
            workbook.close()
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
